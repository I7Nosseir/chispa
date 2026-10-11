"""Builds gulf-leads-and-jobs-2026-10-11.xlsx (+ leads CSV).

Leads: Google Maps via Apify (compass/crawler-google-places), Dubai run on 11 Oct 2026, 80 places.
Chains, franchises, duplicate branches and listings without a phone were removed.
Saudi/Qatar runs could not start: the Apify free credit ran out ($0.45 left). Re-run them next cycle.
Jobs: web search on 11 Oct 2026 (Bayt, Wuzzuf, JobzEG, Himalayas, Indeed, Upwork, Freelancer).
Re-run: python3 commando/leads/build_gulf.py
"""
import csv
from pathlib import Path
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).parent

# (name, niche, area, phone, website_on_google, socials_found, rating, reviews, one_star, special_note)
RAW = [
    ("Essence And Oud Perfumes", "Perfume", "Bur Dubai (Meena Bazar)", "+971 54 783 7846", "", "", 5.0, 17, 0, ""),
    ("FCITY Perfume", "Perfume", "Deira (Gold Souq) + Al Ghubaiba", "+971 55 914 1455", "fcityperfume.com", "", 4.9, 51, 0,
     "Two branches on Maps with only 17 and 34 reviews, and no Instagram/TikTok linked on Google or the site."),
    ("Adaa Lounge Perfumes", "Perfume", "Bur Dubai (Al Fahidi St)", "+971 55 563 6072", "adaaperfumes.com", "", 4.8, 211, 8, ""),
    ("IDOL Perfumes", "Perfume", "Bur Dubai (Meena Bazaar)", "+971 55 849 3585", "idolperfumesdubai.com", "IG+TT", 4.9, 2896, 16,
     "Big review base and active socials already: pitch AI UGC ads + product videos for the e-shop, not a rebuild."),
    ("Mayura Khamas Perfumes", "Perfume", "Bur Dubai (Street 33B)", "+971 4 352 2319", "", "", 4.8, 856, 12,
     "856 reviews at 4.8 and no website or socials on the listing."),
    ("R A M Perfumes", "Perfume", "Deira (Al Souq Al Kabeer)", "+971 52 406 1618", "ramperfumes.com", "IG", 4.8, 90, 3, ""),
    ("Fakhruddin Perfumes (FK Perfumes)", "Perfume", "Deira (Gold Souq)", "+971 4 226 6943", "scentpercent.com", "IG+TT", 5.0, 4007, 26,
     "Maps name, Instagram and website all use different names (Fakhruddin / FK Perfumes / scentpercent.com)."),
    ("Top Naif Abaya Trading", "Abaya", "Deira (Al Manal Centre)", "+971 56 312 1218", "", "", 4.6, 68, 2, ""),
    ("Abaya Fashion (Safi Garden)", "Abaya", "Deira (Souq Al Wasl)", "+971 55 468 1115", "", "", 4.3, 10, 1, ""),
    ("Jawhara Abaya & Sheila", "Abaya", "Deira (Souq Al Wasl)", "+971 55 930 3002", "wa.me link", "", 4.9, 30, 0,
     "Google 'website' is only a WhatsApp link: no catalogue or online store."),
    ("Maitha Abaya", "Abaya", "Deira (Naif Rd)", "+971 55 999 5888", "maithaabaya.com", "IG", 4.8, 300, 12, ""),
    ("Lamya Abayas", "Abaya", "Bur Dubai (50c St)", "+971 58 612 9285", "", "", 4.7, 28, 1, ""),
    ("Qasr Al Ameera Abaya", "Abaya", "Dubai (38th St)", "+971 50 969 6628", "youtube.com channel", "", 4.5, 12, 1,
     "Google 'website' is a YouTube channel."),
    ("Inayah Abaya", "Abaya", "Bur Dubai (Al Ghubaiba bus station)", "+971 54 563 7860", "", "", 5.0, 235, 0,
     "235 reviews at a perfect 5.0 and no website or socials on the listing."),
    ("Abaya Palace", "Abaya", "Bur Dubai (50c St)", "+971 55 567 0370", "instagram.com/abayo.ae", "IG", 4.9, 84, 1,
     "Maps name 'Abaya Palace' but the Instagram is 'abayo.ae'. Two names."),
    ("SAL Dream Gold & Diamonds", "Jewelry", "Karama (Kuwait St)", "+971 4 335 7763", "", "", 5.0, 51, 0, ""),
    ("Mila Gold and Diamonds", "Jewelry", "Karama Center", "+971 4 251 4733", "milagoldanddiamonds.com", "IG+TT", 4.9, 1058, 6, ""),
    ("Nishka Jewellery", "Jewelry", "Karama Center", "+971 56 545 0916", "nishkajewellery.com", "IG", 4.9, 1337, 21, ""),
    ("iSmile Dental Clinic", "Dental", "Dubai (near Jovial Center)", "+971 54 712 8458", "", "", 4.9, 1313, 13,
     "1,313 reviews at 4.9 and no website or socials on the listing."),
    ("National Medical Center", "Clinic", "Jumeirah (Al Diyafah St)", "+971 4 345 3906", "nationalmed.ae", "IG+TT", 4.8, 731, 29, ""),
    ("Real Smile Dental Clinic", "Dental", "Dubai (Al Mina Rd)", "+971 4 355 0112", "facebook.com/realsmiledc", "FB", 5.0, 65, 0, ""),
    ("City Smile Dental Clinic", "Dental", "Bur Dubai (The Atrium)", "+971 4 340 8949", "c-sclinic.com", "IG+TT", 4.8, 214, 8, ""),
    ("DentisTree Dental Clinic", "Dental", "Dubai (Al Mina St)", "+971 4 252 9935", "dentistree.ae", "IG", 4.9, 479, 3, ""),
    ("Pure Smile Dental Clinic", "Dental", "Deira (Doha Center)", "+971 4 575 0768", "puresmiledentalclinic.com", "IG", 5.0, 215, 1, ""),
    ("Al Abrah Clinic", "Dental", "Deira (Al Khor St)", "+971 4 226 3299", "abrahclinics.com", "", 4.8, 196, 6, ""),
    ("The Glow Rituals Aesthetics Clinic", "Aesthetic", "Dubai", "+971 58 871 2188", "", "", 5.0, 128, 0, ""),
    ("Contour Aesthetic Clinic", "Aesthetic", "Dubai (Wasl Port Views)", "+971 4 238 4406", "contourclinicuae.com", "IG", 4.9, 86, 2, ""),
    ("Seagull Aesthetic Clinic", "Aesthetic", "Jumeirah (Al Ghazal Mall)", "+971 4 345 6080", "seagullaesthetic.ae", "IG", 4.6, 211, 18, ""),
    ("Pinas Aesthetic Polyclinic", "Aesthetic", "Dubai (Real Estate Center Bldg)", "+971 4 576 6234", "pinasaesthetic.com", "IG+TT", 4.9, 36, 1, ""),
    ("Cocoage Aesthetics Clinic", "Aesthetic", "Sheikh Zayed Rd (Fairmont Tower)", "+971 54 793 2796", "cocoage.ae", "IG", 4.7, 550, 38, ""),
    ("Halcyon Aesthetics Clinic", "Aesthetic", "Jumeirah Garden", "+971 56 143 0966", "halcyonaestheticsdubai.com", "", 4.9, 86, 2, ""),
    ("Excel Properties", "Real estate", "Sheikh Zayed Rd", "+971 4 336 4447", "excelproperties.ae", "IG+TT", 4.8, 73, 2, ""),
    ("Auric Acres Real Estate", "Real estate", "Bur Dubai (Atrium Center)", "+971 52 957 8280", "auric-acres.com", "", 4.7, 29, 1, ""),
    ("Aalmar Real Estate Broker", "Real estate", "Deira (One Deira Plaza)", "+971 50 276 7550", "aalmargroup.com", "IG", 5.0, 22, 0, ""),
    ("Vihaan Real Estate Brokerage", "Real estate", "Bur Dubai (Al Jawhara 2)", "+971 56 727 6363", "vrebuae.com", "IG", 4.8, 38, 2, ""),
    ("Smart Folks Real Estate Brokers", "Real estate", "Bur Dubai (BurJuman)", "+971 55 282 8039", "themegahub.in", "", 4.8, 52, 2,
     "Google 'website' points to an Indian domain (themegahub.in), not a UAE brokerage site."),
    ("GoldCoast Real Estate Broker", "Real estate", "Umm Hurair (Golden Gate)", "+971 4 399 4944", "goldcoastuae.ae", "IG", 4.8, 40, 2, ""),
    ("Virgo Fitness and Beauty Studio", "Fitness", "Jumeirah (Al Ghazal Mall)", "+971 55 826 0026", "virgodxb.com", "IG", 4.9, 77, 1, ""),
    ("Total Fitness Studio", "Fitness", "Bur Dubai (Al Safa)", "+971 58 548 9016", "", "", 4.5, 35, 4, ""),
    ("Restart Fitness Center & Gym", "Fitness", "Al Mankhool + Al Raffa", "+971 56 891 0323", "restartfitness.ae", "IG", 4.8, 693, 14,
     "The Instagram links on their site are personal coach accounts, not a gym brand account."),
    ("Time 2 Body Fitness Studio", "Fitness", "Al Mankhool Rd", "+971 4 351 0201", "t2bfitness.com", "", 4.5, 45, 3, ""),
    ("AB Fitness", "Fitness", "Dubai", "+971 4 225 6437", "abfitness.ae", "IG+TT", 4.5, 665, 58, ""),
    ("Paulomi's World Of Fitness", "Fitness", "Khalid Bin Al Waleed Rd", "+971 58 551 9113", "", "", 4.9, 906, 20,
     "906 reviews at 4.9 and no website or socials on the listing."),
    ("ANFA Skincare", "Skincare", "Deira (Al Rigga Rd)", "+971 58 895 1408", "anfaskincare.ae", "", 4.9, 411, 2, ""),
    ("Beauty Korea Dubai", "Skincare", "BurJuman (Al Maha Bldg)", "+971 50 969 3232", "beautykoreadubai.com", "IG", 4.6, 285, 13, ""),
    ("Skin Culture Dubai", "Skincare", "Dubai (Crystal Business Center)", "+971 55 655 9858", "skincultureae.com", "", 5.0, 148, 0, ""),
    ("Doppio Roasters", "Coffee", "Dubai (Al Mina St)", "+971 4 324 1686", "doppioroasters.com", "", 4.7, 105, 3, ""),
    ("Specialty Batch Coffee Roastery", "Coffee", "Port Rashid (Marina Cubes)", "+971 4 299 1113", "specialtybatch.com", "IG", 4.7, 43, 0, ""),
    ("mereu. Speciality Coffee", "Coffee", "Dubai (Wasl Port Views)", "+971 55 521 0688", "mereu.ae", "IG+TT", 4.9, 68, 0, ""),
    ("Flat12 Cafe", "Coffee", "Port Rashid (Marina Cubes)", "+971 4 228 7123", "instagram.com/flat12cafe", "IG", 4.7, 1509, 41,
     "1,509 reviews but the Google 'website' is just their Instagram: no menu/order site."),
    ("The Coffee Lab (+ From Scratch)", "Coffee", "Dubai (Marina Views)", "+971 4 546 4670", "thecoffeelab.ae", "IG", 4.7, 4397, 95,
     "From Scratch shares the same phone number on Maps. Two brands, one team: pitch content for both."),
    ("Arabian Fish House Restaurant & Cafe", "Restaurant", "Al Fahidi (Creek St)", "+971 54 993 0057", "arabianfishhouse.com", "IG", 4.9, 2200, 13, ""),
    ("Star Momo House Restaurant", "Restaurant", "Bur Dubai (City Center)", "+971 4 331 2767", "starmomohouse.com", "IG", 4.7, 157, 6, ""),
    ("2 Group Cafe", "Restaurant", "Dubai (Shop W459)", "+971 58 126 4522", "", "", 4.5, 135, 5, ""),
    ("Mehran Restaurant", "Restaurant", "Dubai", "+971 50 346 6061", "", "", 4.0, 77, 7, ""),
    ("New Maritime Restaurant", "Restaurant", "Dubai Maritime City", "+971 56 631 8440", "", "", 3.8, 114, 8, ""),
]

FREE_HOSTS = ("instagram.com", "facebook.com", "youtube.com", "wa.me", "linktr.ee", "themegahub.in")
PRODUCT = {"Perfume", "Abaya", "Jewelry", "Skincare", "Coffee"}
NICHE_BASE = {"Perfume": 7, "Abaya": 6, "Jewelry": 7, "Skincare": 6, "Aesthetic": 7, "Dental": 6, "Clinic": 6,
              "Real estate": 7, "Fitness": 5, "Coffee": 5, "Restaurant": 4}

WHY = {
    "product": "You make AI product shoots and AI UGC ads: new product visuals every week without booking a studio, models or a location. You're also a media buyer, so the creative is made to sell in ads, not just look nice. Price is in USD, below a Dubai studio.",
    "Aesthetic": "Aesthetic clinics need constant before/after-style content and booking ads. You make compliant AI visuals and motion explainers in-house and run the lead campaigns, one person from creative to cost per lead.",
    "Dental": "Clinics win on trust and booking speed. You build the booking site, the content (AI visuals + explainers) and the lead ads, and report cost per booked patient.",
    "Clinic": "Clinics win on trust and booking speed. You build the booking site, the content and the lead ads, and report cost per booked patient.",
    "Real estate": "Brokers live on leads. You've driven 3,000+ qualified leads in travel and know cost-per-lead thinking. You add AI-made property and lifestyle videos so every listing gets an ad creative fast.",
    "Fitness": "Gyms sell on transformation content and offers. You produce reels and AI-made promo visuals and set up a trial sign-up page and ads.",
    "Restaurant": "Restaurants need a steady flow of food video and delivery-app ads. You batch content (real + AI-enhanced) and give them a menu/order site.",
}


def norm_phone(p):
    return p.replace(" ", "")


def score(r):
    name, niche, area, phone, web, soc, rating, reviews, one, note = r
    gaps, need = [], 0
    if note:
        need += 1
        gaps.append(note)
    if not web:
        need += 4
        gaps.append("No website on the Google Maps listing")
    elif any(h in web.lower() for h in FREE_HOSTS):
        need += 3
        gaps.append(f"Google 'website' is not a real site ({web})")
    if not soc:
        need += 2
        gaps.append("No Instagram/TikTok linked on Google or their site")
    elif "TT" not in soc and niche in PRODUCT | {"Aesthetic", "Fitness", "Restaurant"}:
        need += 1
        gaps.append("No TikTok linked, the cheapest place to test UGC-style ads")
    share = one / reviews if reviews else 0
    if share >= 0.10:
        need += 2
        gaps.append(f"{one} one-star of {reviews} ({share:.0%}): reputation work")
    elif share >= 0.06:
        need += 1
        gaps.append(f"{one} one-star of {reviews} ({share:.0%})")
    if rating < 4.0:
        need += 1
        gaps.append(f"Rating {rating}")
    if reviews < 30:
        need += 1
        gaps.append(f"Only {reviews} Google reviews: low visibility")
    if niche in PRODUCT:
        need += 1
        gaps.append("Product business: needs a steady flow of product visuals + ad creatives (AI shoot/UGC fit)")
    need = min(need, 10)
    ability = NICHE_BASE[niche] + (3 if reviews >= 1000 else 2 if reviews >= 300 else 1 if reviews >= 80 else 0)
    return need, min(ability, 10), gaps


def offer(niche, web, soc):
    has_site = web and not any(h in web.lower() for h in FREE_HOSTS)
    parts = []
    if niche in PRODUCT:
        parts.append("AI product shoot (10–20 images) + 3 AI UGC ad videos")
    elif niche in ("Aesthetic", "Dental", "Clinic", "Real estate"):
        parts.append("lead-gen ads + AI ad creatives")
    else:
        parts.append("monthly reels + AI promo visuals")
    if not has_site:
        parts.append("website/e-store with WhatsApp ordering/booking")
    if not soc:
        parts.append("Instagram/TikTok setup + first month of content")
    return "; ".join(parts)


def tier(a):
    return "A: High-paying" if a >= 9 else "B: Mid" if a >= 7 else "C: Low / quick win"


DEAL = {"A: High-paying": "$1,500–3,000 setup + $800–1,500/month",
        "B: Mid": "$400–900 project (AI shoot / UGC pack / landing page) → $500–800/month",
        "C: Low / quick win": "$250–400 starter pack (10 AI product images or 3 UGC ads)"}

leads = []
for r in RAW:
    name, niche, area, phone, web, soc, rating, reviews, one, note = r
    need, ability, gaps = score(r)
    t = tier(ability)
    ph = norm_phone(phone)
    leads.append({
        "Tier": t, "Business": name, "Niche": niche, "City": "Dubai, UAE", "Area": area, "Phone": phone,
        "WhatsApp link (if on WA)": "https://wa.me/" + ph.lstrip("+") if ph.startswith("+9715") else "",
        "Gaps (check before contacting)": " • ".join(gaps), "Offer to pitch": offer(niche, web, soc),
        "Likely deal size (USD)": DEAL[t],
        "Why you (not another agency)": WHY["product"] if niche in PRODUCT else WHY[niche],
        "Rating": rating, "Reviews": reviews, "1★": one, "Website on Google": web or "none", "Socials found": soc or "none",
        "Need": need, "Ability": ability, "Score": need * ability, "Status": "Not contacted", "Notes": "", "Follow-up date": "",
        "Google Maps": "https://www.google.com/maps/search/?api=1&query=" + quote_plus(f"{name} Dubai"),
    })
leads.sort(key=lambda x: (x["Tier"], -x["Score"], -x["Reviews"]))
for i, r in enumerate(leads, 1):
    r["#"] = i
LCOLS = ["#", "Tier", "Business", "Niche", "City", "Area", "Phone", "WhatsApp link (if on WA)", "Gaps (check before contacting)",
         "Offer to pitch", "Likely deal size (USD)", "Why you (not another agency)", "Rating", "Reviews", "1★", "Website on Google",
         "Socials found", "Need", "Ability", "Score", "Status", "Notes", "Follow-up date", "Google Maps"]

# Remote / Gulf jobs found 11 Oct 2026. Verify each is still open before applying.
JOBS = [
    ("1", "Media Buyer", "Malaeb", "Egypt, remote", "Full-time, 2–4 yrs", "Not shown", "Sep 2026",
     "Performance marketing across major platforms, data + AI tools", "Strong fit: lead with Grapes/Rihla numbers",
     "https://www.jobzeg.com/2026/09/media-buyer-remote.html"),
    ("2", "Senior Media Buyer (Remote, Part-time), UAE e-commerce", "Whitecollars", "Remote, Egypt-based candidates only", "Part-time",
     "Not shown", "2026 (check)", "UAE e-commerce ads", "Best fit: part-time fits uni + clients; Gulf e-com experience is the angle",
     "https://himalayas.app/companies/whitecollars/jobs/senior-media-buyer-remote-part-time-uae"),
    ("3", "Media Buyer (mid-career, remote)", "AdFirst", "Egypt, remote", "Full-time", "EGP 24,802–49,603/month", "2026 (check)",
     "Paid media", "Good fit; salary shown is the Egypt benchmark", "https://www.bayt.com/en/egypt/jobs/media-buyer-jobs/"),
    ("4", "Junior Media Buyer", "Attention Theory", "Egypt, remote", "Full-time, 1–2 yrs", "Not shown", "Jun 2026 (may be filled)",
     "Meta Ads + Google Ads", "Apply only if still open; you're above junior",
     "https://www.jobzeg.com/2026/06/junior-media-buyer-job-in-egypt-remote.html"),
    ("5", "Media Buyer (remote, entry level)", "Confidential (Wuzzuf)", "Egypt, remote", "Full-time", "Not shown", "Check on site",
     "Paid media", "Filter Wuzzuf by Remote + last 30 days", "https://wuzzuf.net/a/Media-Buyer-Jobs-in-Egypt"),
    ("6", "E-commerce Specialist (remote)", "Quik Hire Staffing (via JobzEG)", "Egypt, remote", "Full-time, 2+ yrs Shopify/Magento",
     "Claims $60–95K/yr (unverified)", "30 Aug 2026", "Store ops + growth",
     "⚠ Verify the company first; never pay fees or send ID before an interview",
     "https://www.jobzeg.com/2026/08/e-commerce-specialist-remote.html"),
    ("7", "AI Video Creator / Generative Film Specialist", "The Creative Story", "Dubai, in person", "Full-time",
     "AED 12,000–16,000/month", "Oct 2026", "AI video tools, generative film", "Perfect skill match; on-site, so only if you'd relocate. Also pitch them freelance overflow work",
     "https://ae.indeed.com/q-ai-content-creator-jobs.html"),
    ("8", "Growth E-commerce Specialist", "Salla", "Riyadh, on-site", "Full-time", "Not shown", "2026", "E-commerce growth",
     "On-site; useful as a target market: Salla merchants are your Saudi clients", "https://startup.jobs/growth-e-commerce-specialist-riyadh-salla-4545299"),
    ("9", "Performance Marketing Specialist / Senior Media Buyer", "Various (Bayt Saudi)", "Riyadh / Jeddah / Mecca, on-site",
     "Full-time, 2–8 yrs", "Not shown", "Jul–Sep 2026", "TikTok, Snapchat, Google; e-commerce + apps",
     "Shows what Saudi brands want: TikTok + Snapchat. Learn Snap Ads if you haven't", "https://www.bayt.com/en/saudi-arabia/jobs/media-buyer-jobs/"),
    ("10", "Digital media buyer roles across the Middle East", "Various (Bayt)", "Mixed; filter remote", "Mixed", "Some show $500–1,000",
     "Sep 2026", "Paid media", "Filter: Remote. Ignore listings below your floor", "https://www.bayt.com/en/international/jobs/digital-media-buyer-jobs/"),
    ("11", "UGC AI Video Creator (Higgsfield)", "Upwork client", "Worldwide, remote", "Contract-to-hire", "$500 fixed", "May 2026 (closed?)",
     "Higgsfield AI UGC videos", "Proof the gig exists: build an Upwork profile around AI UGC",
     "https://www.upwork.com/freelance-jobs/apply/UGC-Video-Creator-Needed-Higgsfield_~022051153654035466511/"),
    ("12", "Higgsfield AI UGC Video Creator (faceless short-form)", "Upwork client", "Worldwide, remote", "Freelance", "Not shown", "2026",
     "Faceless viral short-form", "Same profile as above", "https://www.upwork.com/freelance-jobs/apply/Higgsfield-UGC-Video-Creator-Needed-Faceless-Content-Viral-Short-Form-Videos_~022064864130430698783/"),
    ("13", "AI UGC Content Creator (Higgsfield / AI video tools)", "Freelancer.com client", "Worldwide, remote", "Freelance",
     "$250–750", "2026", "AI UGC ads", "Quick portfolio builder", "https://www.freelancer.com/projects/ai-content-creation/ugc-content-creator-higgsfield-video"),
    ("14", "Social media / digital marketing remote roles", "Various (Bayt)", "Middle East, remote", "Mixed", "Mixed", "Jul–Sep 2026",
     "Social, content, ads", "Check weekly", "https://www.bayt.com/en/international/jobs/digital-marketing-remote-jobs/"),
]
JCOLS = ["#", "Role", "Company", "Location / who can apply", "Type", "Pay shown", "Posted", "Focus", "Fit + note", "Link",
         "Applied?", "Date", "Result"]

PRICING = [
    ("Job pay (Egypt, local)", "Media buyer in Egypt: job ads estimate EGP 15–25K/month; remote AdFirst role EGP 24.8–49.6K. Survey data (Paylab) is lower, 4.6–17.7K.",
     "Don't go below EGP 25K/month full-time with your results."),
    ("Job pay (Gulf company, remote from Egypt)", "No reliable public data found. Dubai in-person AI video creator: AED 12–16K/month (~$3.3–4.4K).",
     "Estimate: ask $1,000–1,800/month full-time remote, or $15–25/hour part-time. Treat this as a starting point to test."),
    ("AI UGC freelance (global)", "Upwork/Freelancer: $250–750 per project; $500 fixed seen for a Higgsfield UGC set.",
     "3-video AI UGC pack: $300–500. 10-video monthly: $900–1,500."),
    ("Product photos (benchmark)", "Studio packshot $25–50/image; styled lifestyle $150–500+/image; half-day shoot $500–2,500 (2026 guides). AI is quoted 80–90% cheaper.",
     "AI product shoot: 10 images $250–400; 20 images + 2 short videos $600–900. Still cheaper than a studio."),
    ("Websites (Gulf)", "Your profile: websites from ~$600 for Gulf clients (proposed).",
     "Landing page $600–900; business site $1,200–2,000; e-store (Salla/Shopify setup + design) $1,500–3,000."),
    ("Monthly retainer (Gulf)", "Your profile: social from ~$400/month (proposed).",
     "Content + AI creatives $500–900/month; with ads management add $400+ or 10–12% of spend."),
]


def style(ws, widths, tier_col=None):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if tier_col is not None:
            row[tier_col].fill = PatternFill("solid", fgColor={"A": "D1FAE5", "B": "FEF3C7", "C": "F3F4F6"}[row[tier_col].value[0]])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
ws = wb.active
ws.title = "Gulf leads (Dubai)"
ws.append(LCOLS)
for r in leads:
    ws.append([r[c] for c in LCOLS])
style(ws, [5, 16, 32, 11, 11, 26, 17, 26, 70, 42, 30, 60, 7, 8, 6, 26, 10, 6, 7, 7, 14, 30, 12, 30], tier_col=1)
dv = DataValidation(type="list", formula1='"Not contacted,WhatsApp sent,DM sent,Email sent,Replied,Sample sent,Call booked,Won,Not interested"')
ws.add_data_validation(dv)
col = get_column_letter(LCOLS.index("Status") + 1)
dv.add(f"{col}2:{col}{len(leads) + 1}")

j = wb.create_sheet("Remote jobs")
j.append(JCOLS)
for row in JOBS:
    j.append(list(row) + ["", "", ""])
style(j, [4, 34, 24, 26, 18, 22, 16, 30, 46, 50, 10, 12, 16])

p = wb.create_sheet("Fair pricing")
p.append(["What", "Market data found (11 Oct 2026)", "Your ask (proposed: confirm before quoting)"])
for row in PRICING:
    p.append(list(row))
style(p, [30, 80, 80])

w = wb.create_sheet("Where to look weekly")
w.append(["Channel", "Search / filter", "Why"])
for row in [
    ("LinkedIn Jobs", "\"media buyer\" OR \"performance marketing\" OR \"AI content\", Remote, past week; locations Egypt, UAE, Saudi Arabia", "Most Gulf brands post here first"),
    ("Wuzzuf", "Media Buyer / Media Buying / TikTok Ads, Remote, last 30 days", "Egypt remote roles"),
    ("Bayt", "Digital media buyer + Remote filter, Middle East", "Gulf roles, some remote"),
    ("Workable", "jobs.workable.com, Egypt, media buyer", "Startups hiring remote"),
    ("Upwork", "AI UGC, Higgsfield, AI product photography, Meta ads Arabic", "Global freelance gigs"),
    ("Mostaql (مستقل) / Khamsat", "إعلانات ممولة، تصميم منتجات بالذكاء الاصطناعي، متجر سلة", "Saudi/Gulf clients hiring in Arabic"),
    ("Contra / Wellfound", "AI content, creative strategist, remote", "Startups that pay in USD"),
]:
    w.append(list(row))
style(w, [26, 80, 40])

wb.save(OUT / "gulf-leads-and-jobs-2026-10-11.xlsx")
with open(OUT / "gulf-leads-2026-10-11.csv", "w", newline="", encoding="utf-8-sig") as f:
    wr = csv.DictWriter(f, fieldnames=LCOLS)
    wr.writeheader()
    for r in leads:
        wr.writerow({c: r[c] for c in LCOLS})
print(len(leads), "leads", {t: sum(1 for r in leads if r["Tier"] == t) for t in DEAL})
for r in leads[:12]:
    print(r["#"], r["Tier"][0], r["Business"], r["Score"])
